from pathlib import Path
from openpyxl import load_workbook
from _poc_helpers import setup, write_csv, validate, run, read_json


def test_reproducibility(tmp_path, monkeypatch):
    roots=[]; results=[]
    for n in ("a","b"):
        root=tmp_path/n; root.mkdir(); c=setup(root,monkeypatch)
        for i in range(10): write_csv(Path("input")/f"file_{i+1:02}.csv")
        results.append((run([validate(p,c) for p in sorted(Path("input").iterdir())],c), c))
    a=load_workbook(results[0][0]["output_path"],data_only=True); b=load_workbook(results[1][0]["output_path"],data_only=True)
    for sheet in ("Data","Summary"):
        av=list(a[sheet].values); bv=list(b[sheet].values)
        if sheet=="Summary":
            av[1][-1]=bv[1][-1]=None
        assert av==bv
    ma=read_json(results[0][0]["manifest_path"]); mb=read_json(results[1][0]["manifest_path"])
    for k in ("status","rows_received","rows_accepted","rows_rejected"): assert ma[k]==mb[k]
