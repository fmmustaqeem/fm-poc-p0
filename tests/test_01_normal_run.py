from pathlib import Path
from openpyxl import load_workbook
from _poc_helpers import setup, write_csv, validate, run


def test_normal_run(tmp_path, monkeypatch):
    c = setup(tmp_path, monkeypatch)
    files = []
    for i in range(10):
        p = Path("input") / f"file_{i+1:02}.csv"; write_csv(p); files.append(validate(p, c))
    r = run(files, c); m = __import__("_poc_helpers").read_json(r["manifest_path"])
    assert m["rows_received"] == 50 and m["rows_accepted"] == 50 and m["rows_rejected"] == 0
    assert m["status"] == "success"
    wb = load_workbook(r["output_path"]); assert wb["Data"].max_row == 51
    assert not list(Path("quarantine").iterdir())
